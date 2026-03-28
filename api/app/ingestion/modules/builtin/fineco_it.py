"""
FinecoIT - FinecoBank Italy bank statement extractor.
Supported formats: XLSX (Excel)
Export path in Fineco UI: Portafoglio → Movimenti → Esporta → Excel
Sample FINGERPRINT keywords (from actual file): "conto corrente", "data_operazione", "entrate", "uscite", "descrizione_completa"

The export begins with several header rows including account info, followed by a notes section,
then a "Risultati Ricerca" label, and finally the column header row at row index 12:
  Data_Operazione | Data_Valuta | Entrate | Uscite | Descrizione | Descrizione_Completa | Stato | Moneymap

Each transaction row may have a credit (Entrate > 0) or a debit (Uscite < 0, stored as negative float).
"""

from __future__ import annotations

from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl

FINGERPRINT: list[str] = [
    "conto corrente",
    "data_operazione",
    "entrate",
    "uscite",
    "descrizione_completa",
    "risultati ricerca",
]
MODULE_VERSION: str = "1.0.0"

# Header names expected in the data header row (lowercased)
_DATE_OP = "data_operazione"
_DATE_VAL = "data_valuta"
_CREDIT = "entrate"
_DEBIT = "uscite"
_DESC = "descrizione"
_DESC_FULL = "descrizione_completa"
_STATE = "stato"


def extract(file_path: str | Path):  # noqa: ANN201  (return type is ExtractionResult)
    """Extract transactions from FinecoBank XLSX export."""
    from app.schemas.ingestion import ExtractionResult, ExtractedDocument, Transaction  # noqa: PLC0415

    path = Path(file_path)
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    # ---- locate the header row ----
    header_idx: int | None = None
    headers: list[str] = []
    for i, row in enumerate(rows[:30]):
        row_strs = [str(c).lower().strip() if c is not None else "" for c in row]
        if _DATE_OP in row_strs:
            header_idx = i
            headers = row_strs
            break

    if header_idx is None:
        return ExtractionResult(
            document=ExtractedDocument(
                document_type="bank_statement",
                module_name="FinecoIT",
                module_source="builtin",
                module_version=MODULE_VERSION,
            ),
            confidence=0.3,
            tier_used="module",
            warnings=["Could not detect FinecoIT header row"],
        )

    # ---- map column indices ----
    def _col(name: str) -> int | None:
        try:
            return headers.index(name)
        except ValueError:
            return None

    date_col = _col(_DATE_OP)
    credit_col = _col(_CREDIT)
    debit_col = _col(_DEBIT)
    desc_col = _col(_DESC_FULL) if _col(_DESC_FULL) is not None else _col(_DESC)
    state_col = _col(_STATE)

    if date_col is None:
        return ExtractionResult(
            document=ExtractedDocument(
                document_type="bank_statement",
                module_name="FinecoIT",
                module_source="builtin",
                module_version=MODULE_VERSION,
            ),
            confidence=0.3,
            tier_used="module",
            warnings=["FinecoIT: date column not found in header"],
        )

    # ---- extract metadata from top rows ----
    account_holder: str | None = None
    statement_start = None
    statement_end = None
    for row in rows[:10]:
        cell0 = str(row[0]).strip() if row[0] is not None else ""
        if "intestazione conto corrente:" in cell0.lower():
            account_holder = cell0.split(":", 1)[-1].strip()
        if "periodo dal:" in cell0.lower():
            # "Periodo Dal: 24/03/2025 Al: 24/03/2026"
            import re  # noqa: PLC0415

            m = re.search(
                r"dal:\s*(\d{2}/\d{2}/\d{4}).*al:\s*(\d{2}/\d{2}/\d{4})",
                cell0,
                re.IGNORECASE,
            )
            if m:
                from datetime import datetime  # noqa: PLC0415

                try:
                    statement_start = datetime.strptime(m.group(1), "%d/%m/%Y").date()
                    statement_end = datetime.strptime(m.group(2), "%d/%m/%Y").date()
                except ValueError:
                    pass

    # ---- parse transactions ----
    transactions: list[Transaction] = []
    warnings: list[str] = []

    for row in rows[header_idx + 1 :]:
        if not row or all(c is None for c in row):
            continue
        # skip "Contabilizzato" filter - accept all states including reverted ones
        try:
            date_val = row[date_col] if date_col is not None else None
            if date_val is None:
                continue

            # date may be a datetime object or a string
            if hasattr(date_val, "date"):
                txn_date = date_val.date()
            elif isinstance(date_val, str):
                from datetime import datetime  # noqa: PLC0415

                for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
                    try:
                        txn_date = datetime.strptime(date_val.strip(), fmt).date()
                        break
                    except ValueError:
                        continue
                else:
                    continue
            else:
                continue

            # amount: Entrate (credit, positive) or Uscite (debit, negative)
            credit_raw = (
                row[credit_col]
                if credit_col is not None and credit_col < len(row)
                else None
            )
            debit_raw = (
                row[debit_col]
                if debit_col is not None and debit_col < len(row)
                else None
            )

            amount: Decimal | None = None
            if credit_raw not in (None, "", 0):
                amount = Decimal(str(credit_raw)).quantize(Decimal("0.0001"))
            elif debit_raw not in (None, "", 0):
                # Uscite is stored as negative float in the XLSX
                amount = Decimal(str(debit_raw)).quantize(Decimal("0.0001"))
            else:
                continue  # row has no amount - skip

            desc_raw = (
                row[desc_col] if desc_col is not None and desc_col < len(row) else None
            )
            description = str(desc_raw).strip() if desc_raw is not None else ""

            transactions.append(
                Transaction(
                    date=txn_date,
                    description=description,
                    amount=amount,
                    currency="EUR",
                )
            )
        except (InvalidOperation, IndexError, TypeError):
            continue

    return ExtractionResult(
        document=ExtractedDocument(
            document_type="bank_statement",
            module_name="FinecoIT",
            module_source="builtin",
            module_version=MODULE_VERSION,
            transactions=transactions,
            account_holder=account_holder,
            statement_period_start=statement_start,
            statement_period_end=statement_end,
        ),
        confidence=0.9 if transactions else 0.4,
        tier_used="module",
        warnings=warnings
        if warnings
        else (["No transactions extracted"] if not transactions else []),
    )
